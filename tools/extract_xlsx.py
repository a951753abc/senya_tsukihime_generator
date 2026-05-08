"""
從 Excel 角色卡（charcard.xlsx）抽 reference 資料：
  - 風格（styles）→ data/styles.json
  - 感情（emotions）→ data/emotions.json

Usage:
    python extract_xlsx.py <xlsx_path> <out_dir>
"""
from __future__ import annotations
import sys
import json
import re
from pathlib import Path
from openpyxl import load_workbook


def _clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    if s in ("-", "—", "－"):
        return ""
    return s


def _infer_category(name: str, class_restriction: str) -> str:
    """依名稱與級別限制推導下拉群組分類"""
    if name.startswith("起源"):
        return "起源系"
    if class_restriction and class_restriction not in ("", "無"):
        return "限定系"
    return "通用系"


def extract_styles(wb) -> list[dict]:
    if "風格" not in wb.sheetnames:
        return []
    ws = wb["風格"]
    styles = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue  # header
        name = _clean(row[0])
        if not name:
            continue
        description = _clean(row[1])
        antithesis_raw = _clean(row[2])
        # 背反律常以「...」包住，去除
        antithesis = re.sub(r"^「(.+)」$", r"\1", antithesis_raw)
        classification = _clean(row[3])
        compensation = _clean(row[4])
        class_restriction = _clean(row[5])
        emotion = _clean(row[6])
        representative = _clean(row[7]) if len(row) > 7 else ""
        styles.append({
            "name": name,
            "description": description,
            "classification": classification,
            "antithesis": antithesis,
            "compensation": compensation,
            "classRestriction": class_restriction,
            "emotion": emotion,
            "representative": representative,
            "category": _infer_category(name, class_restriction),
        })
    return styles


def extract_emotions(wb) -> list[dict]:
    """感情工作表 — 表頭可能是 (感情, 羈絆值, 色調, 備註)，含分組標題行"""
    if "感情" not in wb.sheetnames:
        return []
    ws = wb["感情"]
    emotions = []
    current_section = ""
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        col0 = _clean(row[0])
        col1 = row[1]
        col2 = _clean(row[2])
        col3 = _clean(row[3]) if len(row) > 3 else ""
        if not col0:
            continue
        # 分組標題（如「基本感情 ▲ 創角時所獲得並使用的感情」）
        if "▲" in col0 and (col1 is None or _clean(col1) == ""):
            current_section = re.sub(r"\s*▲.*$", "", col0).strip()
            continue
        # 一般資料列
        try:
            bond = int(float(col1)) if col1 is not None else 0
        except (ValueError, TypeError):
            bond = 0
        tone_raw = col2
        if tone_raw == "潔淨":
            tone = "pure"
        elif tone_raw == "瘋狂":
            tone = "crazy"
        else:
            tone = ""
        emotions.append({
            "name": col0,
            "bond": bond,
            "tone": tone,
            "section": current_section,
            "note": col3,
        })
    return emotions


def main():
    if len(sys.argv) < 3:
        print("Usage: python extract_xlsx.py <xlsx_path> <out_dir>")
        sys.exit(1)
    sys.stdout.reconfigure(encoding="utf-8")
    xlsx_path = Path(sys.argv[1])
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(xlsx_path, data_only=True)

    styles = extract_styles(wb)
    (out_dir / "styles.json").write_text(
        json.dumps({"schemaVersion": "2.0", "styles": styles}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"OK styles    {len(styles):3d} entries → styles.json")

    emotions = extract_emotions(wb)
    (out_dir / "emotions.json").write_text(
        json.dumps({"schemaVersion": "2.0", "emotions": emotions}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"OK emotions  {len(emotions):3d} entries → emotions.json")

    # 統計
    by_cat = {}
    for s in styles:
        by_cat[s["category"]] = by_cat.get(s["category"], 0) + 1
    print(f"   styles 分類: {by_cat}")
    by_tone = {}
    for e in emotions:
        by_tone[e["tone"]] = by_tone.get(e["tone"], 0) + 1
    print(f"   emotions 色調: {by_tone}")


if __name__ == "__main__":
    main()
